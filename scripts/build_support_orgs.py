#!/usr/bin/env python3
"""出入国在留管理庁「登録支援機関登録簿」(Excel) から data/support-orgs.json を生成する。

使い方:
    python3 scripts/build_support_orgs.py            # 掲載ページから最新のExcelを取得して生成
    python3 scripts/build_support_orgs.py FILE.xlsx  # 手元のExcelから生成

要件: openpyxl（pip install openpyxl）と curl
ライセンス: 出入国在留管理庁ホームページのコンテンツは公共データ利用規約（PDL1.0）が適用される。
代表者氏名はツールの用途に不要なため取り込まない（データ最小化）。
"""
import datetime
import json
import re
import subprocess
import sys
import tempfile
import unicodedata
from pathlib import Path

import openpyxl

PAGE_URL = "https://www.moj.go.jp/isa/applications/ssw/nyuukokukanri07_00205.html"
ORIGIN = "https://www.moj.go.jp"
# 既定のUAだとCloudFrontに403で拒否される
UA = "Mozilla/5.0 (compatible; mcp-tools-data-builder)"
OUT = Path(__file__).resolve().parent.parent / "data" / "support-orgs.json"
HEADER_ROWS = 6
MIN_ORGS = 5000


def get(url: str) -> bytes:
    # Python の urllib は接続を切られる（2026-10-13 実測。curl は通る）ため curl で取得する
    return subprocess.run(
        ["curl", "-sSfL", "--max-time", "120", "-A", UA, url], check=True, capture_output=True
    ).stdout


def latest_xlsx_url() -> str:
    html = get(PAGE_URL).decode("utf-8", "replace")
    m = re.search(r'href="([^"]+\.xlsx)"[^>]*>\s*登録支援機関登録簿', html)
    if not m:
        sys.exit("掲載ページに登録支援機関登録簿のExcelリンクが見つかりません")
    return ORIGIN + m.group(1)


def text(v) -> str:
    if v is None:
        return ""
    s = unicodedata.normalize("NFKC", str(v))
    return re.sub(r"\s+", " ", s).strip()


def date(v) -> str:
    if hasattr(v, "strftime"):
        return v.strftime("%Y-%m-%d")
    # 日付書式の無いセルは Excel のシリアル値（1899-12-30 起点の日数）で入っている（2026-10-08版で63件）
    if isinstance(v, (int, float)) or re.fullmatch(r"\d{5}", text(v)):
        return (datetime.date(1899, 12, 30) + datetime.timedelta(days=int(float(v)))).isoformat()
    s = text(v)
    m = re.fullmatch(r"(\d{4})[-/.年](\d{1,2})[-/.月](\d{1,2})日?", s)
    if not m:
        sys.exit(f"登録日を解釈できません: {s!r}")
    return f"{m.group(1)}-{int(m.group(2)):02d}-{int(m.group(3)):02d}"


def as_of(ws) -> str:
    for row in ws.iter_rows(min_row=1, max_row=2, values_only=True, max_col=20):
        for v in row:
            m = re.search(r"(\d{4})年(\d{1,2})月(\d{1,2})日現在", text(v))
            if m:
                return f"{m.group(1)}-{int(m.group(2)):02d}-{int(m.group(3)):02d}"
    sys.exit("基準日（○年○月○日現在）が見つかりません")


def build(path: str) -> dict:
    ws = openpyxl.load_workbook(path, read_only=True).worksheets[0]
    header = [text(v) for v in next(ws.iter_rows(min_row=3, max_row=3, values_only=True, max_col=14))]
    # 列位置の前提が崩れたら黙って誤データを作らずに止める
    expected = {0: "登録番号", 2: "氏名又は名称", 3: "住所", 7: "支援を行う事務所", 10: "支援業務の内容", 12: "相談に応じる体制", 13: "備考"}
    for i, word in expected.items():
        if word not in header[i]:
            sys.exit(f"列{i}の見出しが想定と異なります: {header[i]!r}（期待: {word}）")
    orgs = []
    for r in ws.iter_rows(min_row=HEADER_ROWS + 1, values_only=True, max_col=14):
        if not r[0]:
            continue
        support = text(r[10])
        orgs.append(
            {
                "registrationNumber": text(r[0]),
                "registeredOn": date(r[1]),
                "name": text(r[2]),
                "postCode": text(r[3]),
                "address": text(r[4]),
                "phone": text(r[5]),
                "officeName": text(r[7]),
                "officeAddress": text(r[9]),
                "languages": text(r[12]).replace("(対応可能言語)", "").strip(),
                "optionalSupport": "☑ 上記に加え" in support,
                "remarks": text(r[13]),
            }
        )
    return {"asOf": as_of(ws), "sourceUrl": PAGE_URL, "organizations": orgs}


def main() -> None:
    if len(sys.argv) > 1:
        data = build(sys.argv[1])
    else:
        url = latest_xlsx_url()
        print(f"download: {url}", file=sys.stderr)
        with tempfile.NamedTemporaryFile(suffix=".xlsx") as tmp:
            tmp.write(get(url))
            tmp.flush()
            data = build(tmp.name)
    # 2026-10-08 時点で 11,660件。桁違いに少なければ取得か解析の失敗とみなし、上書きしない
    if len(data["organizations"]) < MIN_ORGS:
        sys.exit(f"件数が少なすぎます（{len(data['organizations'])}件 < {MIN_ORGS}件）。出力しません")
    OUT.parent.mkdir(exist_ok=True)
    OUT.write_text(json.dumps(data, ensure_ascii=False, separators=(",", ":")) + "\n", encoding="utf-8")
    print(f"wrote {OUT} ({len(data['organizations'])} orgs, asOf {data['asOf']})", file=sys.stderr)


if __name__ == "__main__":
    main()
